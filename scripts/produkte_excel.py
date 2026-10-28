"""Produkte als Excel-Liste austauschen (vor dem Livegang, solange das CMS noch nicht angeschlossen ist).

    python3 scripts/produkte_excel.py export produkte.xlsx   # alle Produkte in eine Excel-Datei
    python3 scripts/produkte_excel.py import produkte.xlsx   # bearbeitete Datei zurück in src/content/produkte

Beim Import werden nur Produkte neu geschrieben, deren Inhalt sich geändert hat.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ORDNER = Path('src/content/produkte')
GRUPPEN = json.loads(Path('src/data/produktgruppen.json').read_text(encoding='utf-8'))['gruppen']
GRUPPE_TITEL = {g['slug']: g['titel'] for g in GRUPPEN}
GRUPPE_SLUG = {g['titel']: g['slug'] for g in GRUPPEN}
MARKE_TITEL = {None: 'wie Gruppe', 'tafag': 'TAFAG'}
MARKE_WERT = {v: k for k, v in MARKE_TITEL.items()}
STATUS_TITEL = {'aktiv': 'aktiv', 'abgekuendigt': 'abgekündigt'}
STATUS_WERT = {v: k for k, v in STATUS_TITEL.items()}

# Spalten im Blatt «Produkte»: (Überschrift, Feld, Breite)
SPALTEN = [
    ('ID (nicht ändern)', 'id', 30),
    ('Aktion', 'aktion', 11),
    ('Titel', 'titel', 30),
    ('Gruppe', 'gruppe', 32),
    ('Marke', 'marke', 12),
    ('Kategorie (Filter)', 'kategorie', 18),
    ('Kurztext', 'kurz', 50),
    ('Beschreibung', 'beschreibung', 80),
    ('Status', 'status', 13),
    ('Hinweis', 'hinweis', 40),
    ('Reihenfolge', 'reihenfolge', 12),
    ('Bild', 'bild', 40),
    ('Entwurf', 'entwurf', 10),
]
# Unterlisten: Blattname → (Feld im Produkt, [(Überschrift, Unterfeld, Breite)])
LISTEN = {
    'Technische Daten': ('technische_daten', [('Merkmal', 'name', 30), ('Wert', 'wert', 60)]),
    'Modelle': ('modelle', [('Name', 'name', 30), ('Text', 'text', 70)]),
    'Downloads': ('downloads', [('Datei', 'datei', 60), ('Bezeichnung', 'bezeichnung', 40)]),
}
REIHENFOLGE_FELDER = ['titel', 'gruppe', 'marke', 'kategorie', 'kurz', 'bild', 'beschreibung', 'technische_daten',
                      'modelle', 'downloads', 'reihenfolge', 'status', 'hinweis', 'entwurf']

SCHRIFT = 'Arial'
KOPF = PatternFill('solid', fgColor='0B2239')
GESPERRT = PatternFill('solid', fgColor='E8EEF5')


def lade_produkte():
    produkte = {}
    for f in sorted(ORDNER.glob('*.yml')):
        produkte[f.stem] = yaml.safe_load(f.read_text(encoding='utf-8')) or {}
    return produkte


def kopfzeile(ws, spalten):
    for i, (titel, _, breite) in enumerate(spalten, start=1):
        c = ws.cell(row=1, column=i, value=titel)
        c.font = Font(name=SCHRIFT, bold=True, color='FFFFFF')
        c.fill = KOPF
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = breite
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions


def zelle(ws, r, c, wert, gesperrt=False):
    z = ws.cell(row=r, column=c, value=wert)
    z.font = Font(name=SCHRIFT)
    z.alignment = Alignment(vertical='top', wrap_text=True)
    if gesperrt:
        z.fill = GESPERRT
    return z


def auswahl(ws, spalte, werte, zeilen=500):
    dv = DataValidation(type='list', formula1='"' + ','.join(werte) + '"', allow_blank=True)
    dv.error = 'Bitte einen Wert aus der Liste wählen.'
    ws.add_data_validation(dv)
    dv.add(f'{spalte}2:{spalte}{zeilen}')


def export(ziel):
    produkte = lade_produkte()
    wb = Workbook()

    an = wb.active
    an.title = 'Anleitung'
    an.column_dimensions['A'].width = 110
    zeilen = [
        ('Produkte der Website – zum Prüfen, Ändern und Neu-Anlegen', True),
        ('', False),
        ('Blatt «Produkte»: eine Zeile pro Produkt. Texte direkt in den Zellen ändern.', False),
        ('Spalte «ID» (grau) nicht ändern – sie verbindet die Zeile mit dem Produkt auf der Website.', False),
        ('Neues Produkt: neue Zeile, «ID» leer lassen, bei «Aktion» «neu» wählen, mindestens Titel und Gruppe ausfüllen.', False),
        ('Produkt entfernen: Zeile stehen lassen und bei «Aktion» «löschen» wählen.', False),
        ('Marke: «wie Gruppe» = Marke der Produktgruppe (TAFAG hat zurzeit nur eine Marke).', False),
        ('Beschreibung: Absätze mit Leerzeile trennen; **fett** und *kursiv* sind möglich.', False),
        ('', False),
        ('Blätter «Technische Daten», «Modelle», «Downloads»: eine Zeile pro Eintrag, Spalte A = ID des Produkts.', False),
        ('Bei neuen Produkten in Spalte A den Titel des neuen Produkts eintragen.', False),
        ('Reihenfolge der Zeilen = Reihenfolge auf der Website.', False),
        ('', False),
        ('Bilder und Datenblätter: als Dateien mitschicken, benannt nach dem Produkt (z. B. zagro-e-maxi-m.jpg, vtms-datenblatt.pdf).', False),
        ('In den Spalten «Bild» bzw. «Datei» genügt dann der Dateiname.', False),
        ('', False),
        ('Beispiel für eine neue Zeile im Blatt «Produkte»:', True),
        ('ID: (leer) | Aktion: neu | Titel: Kartonverschliesser XY | Gruppe: Maschinen, Fördertechnik & Intralogistik | '
         'Marke: wie Gruppe | Kategorie: Kartonverschliessmaschinen | Kurztext: Ein Satz für die Übersicht | Status: aktiv', False),
    ]
    for i, (text, fett) in enumerate(zeilen, start=1):
        c = an.cell(row=i, column=1, value=text)
        c.font = Font(name=SCHRIFT, bold=fett, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)

    ws = wb.create_sheet('Produkte')
    kopfzeile(ws, SPALTEN)
    ids = sorted(produkte, key=lambda k: (list(GRUPPE_TITEL).index(produkte[k].get('gruppe')), produkte[k].get('reihenfolge', 100), k))
    for r, pid in enumerate(ids, start=2):
        p = produkte[pid]
        werte = {
            'id': pid, 'aktion': None, 'titel': p.get('titel'), 'gruppe': GRUPPE_TITEL.get(p.get('gruppe')),
            'marke': MARKE_TITEL[p.get('marke')], 'kategorie': p.get('kategorie'), 'kurz': p.get('kurz'),
            'beschreibung': p.get('beschreibung'), 'status': STATUS_TITEL[p.get('status', 'aktiv')],
            'hinweis': p.get('hinweis'), 'reihenfolge': p.get('reihenfolge'), 'bild': p.get('bild'),
            'entwurf': 'ja' if p.get('entwurf') else 'nein',
        }
        for c, (_, feld, _) in enumerate(SPALTEN, start=1):
            zelle(ws, r, c, werte[feld], gesperrt=(feld == 'id'))
    spalte = {feld: ws.cell(row=1, column=i).column_letter for i, (_, feld, _) in enumerate(SPALTEN, start=1)}
    auswahl(ws, spalte['aktion'], ['neu', 'löschen'])
    auswahl(ws, spalte['gruppe'], list(GRUPPE_SLUG))
    auswahl(ws, spalte['marke'], list(MARKE_WERT))
    auswahl(ws, spalte['status'], list(STATUS_WERT))
    auswahl(ws, spalte['entwurf'], ['ja', 'nein'])

    for blatt, (feld, unter) in LISTEN.items():
        wl = wb.create_sheet(blatt)
        kopfzeile(wl, [('ID (Produkt)', 'id', 30)] + unter)
        r = 2
        for pid in ids:
            for eintrag in produkte[pid].get(feld) or []:
                zelle(wl, r, 1, pid, gesperrt=True)
                for c, (_, unterfeld, _) in enumerate(unter, start=2):
                    zelle(wl, r, c, eintrag.get(unterfeld))
                r += 1

    wb.active = 1
    wb.save(ziel)
    print(f'{len(ids)} Produkte exportiert nach {ziel}')


def slug(text):
    t = unicodedata.normalize('NFKD', text.replace('ä', 'ae').replace('ö', 'oe').replace('ü', 'ue')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', t.lower()).strip('-')


def text(v):
    if v is None:
        return None
    v = str(v).strip()
    return v or None


class Dumper(yaml.SafeDumper):
    pass


def _str(d, s):
    # Mehrzeilige Texte als Block, damit sie im Repository lesbar bleiben
    return d.represent_scalar('tag:yaml.org,2002:str', s, style='|' if '\n' in s else None)


Dumper.add_representer(str, _str)


def importieren(quelle):
    wb = load_workbook(quelle, data_only=True)
    alt = lade_produkte()
    ws = wb['Produkte']
    kopf = [c.value for c in ws[1]]
    feld_von = {titel: feld for titel, feld, _ in SPALTEN}
    neu, titel_zu_id, loeschen = {}, {}, []
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        w = {feld_von[k]: v for k, v in zip(kopf, zeile) if k in feld_von}
        if not any(text(v) for v in w.values()):
            continue
        pid, aktion = text(w.get('id')), text(w.get('aktion'))
        if aktion == 'löschen':
            if pid:
                loeschen.append(pid)
            continue
        if not text(w.get('titel')) or text(w.get('gruppe')) not in GRUPPE_SLUG:
            sys.exit(f'Zeile ohne Titel oder mit unbekannter Gruppe: {w}')
        if not pid:
            pid = slug(w['titel'])
            if pid in alt or pid in neu:
                sys.exit(f'ID «{pid}» gibt es schon – bitte Titel eindeutig machen.')
        titel_zu_id[text(w['titel'])] = pid
        p = {
            'titel': text(w['titel']),
            'gruppe': GRUPPE_SLUG[text(w['gruppe'])],
            'marke': MARKE_WERT.get(text(w.get('marke')) or 'wie Gruppe'),
            'kategorie': text(w.get('kategorie')),
            'kurz': text(w.get('kurz')),
            'bild': text(w.get('bild')),
            'beschreibung': text(w.get('beschreibung')),
            'reihenfolge': int(w['reihenfolge']) if text(w.get('reihenfolge')) else None,
            'status': None if STATUS_WERT.get(text(w.get('status')) or 'aktiv') == 'aktiv' else 'abgekuendigt',
            'hinweis': text(w.get('hinweis')),
            'entwurf': True if text(w.get('entwurf')) == 'ja' else None,
        }
        # Bildpfad: blosser Dateiname → Ablageort der Uploads
        if p['bild'] and '/' not in p['bild']:
            p['bild'] = '/uploads/produkte/' + p['bild']
        neu[pid] = p

    for blatt, (feld, unter) in LISTEN.items():
        wl = wb[blatt]
        for zeile in wl.iter_rows(min_row=2, values_only=True):
            pid = text(zeile[0])
            if not pid:
                continue
            pid = titel_zu_id.get(pid, pid)
            if pid not in neu:
                continue
            eintrag = {uf: text(v) for (_, uf, _), v in zip(unter, zeile[1:])}
            if feld == 'downloads' and eintrag.get('datei') and '/' not in eintrag['datei']:
                eintrag['datei'] = '/uploads/downloads/produkte/' + eintrag['datei']
            eintrag = {k: v for k, v in eintrag.items() if v is not None}
            if eintrag:
                neu[pid].setdefault(feld, []).append(eintrag)

    geaendert, angelegt = [], []
    for pid, p in neu.items():
        daten = {k: p[k] for k in REIHENFOLGE_FELDER if p.get(k) not in (None, [])}
        vorher = {k: v for k, v in alt.get(pid, {}).items() if v not in (None, [], '') and not (k == 'status' and v == 'aktiv')}
        if daten == vorher:
            continue
        (ORDNER / f'{pid}.yml').write_text(
            yaml.dump(daten, Dumper=Dumper, allow_unicode=True, sort_keys=False, width=1000), encoding='utf-8')
        (geaendert if pid in alt else angelegt).append(pid)
    for pid in loeschen:
        if (ORDNER / f'{pid}.yml').exists():
            (ORDNER / f'{pid}.yml').unlink()
    print('Geändert:', ', '.join(geaendert) or '–')
    print('Neu:', ', '.join(angelegt) or '–')
    print('Gelöscht:', ', '.join(loeschen) or '–')


if __name__ == '__main__':
    if len(sys.argv) != 3 or sys.argv[1] not in ('export', 'import'):
        sys.exit(__doc__)
    export(sys.argv[2]) if sys.argv[1] == 'export' else importieren(sys.argv[2])
